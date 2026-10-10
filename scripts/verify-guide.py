"""Check the rendered manual against its authored index and figure inventory."""
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from pypdf import PdfReader
import pdfplumber

root = Path(__file__).resolve().parents[1]
directory = root / 'artifacts/guide'
layout = json.loads((directory / 'guide-layout.json').read_text(encoding='utf8'))
file = directory / f"EluShell操作指南-{layout['version']}.docx"
reader = PdfReader(directory / 'rendered/guide.pdf')
document = Document(file)
norm = lambda s: re.sub(r'\s+', '', s)
assert len(reader.pages) == layout['pages'], (len(reader.pages), layout['pages'])
geometric = pdfplumber.open(directory / 'rendered/guide.pdf')
for number, (title, page) in enumerate(zip(layout['titles'], reader.pages), 1):
    # LibreOffice emits fallback CJK fonts in separate runs. Sort heading glyphs
    # by physical position rather than the PDF stream's font-run order.
    heading = ''.join(c['text'] for c in sorted(geometric.pages[number-1].crop((40,40,580,85)).chars, key=lambda c:c['x0']))
    assert norm(title) in norm(page.extract_text()) or norm(title) in norm(heading), (number, title, heading)
    assert ('EluShell'+layout['version']) in norm(page.extract_text()), ('footer', number)
index = next(table for table in document.tables if table.cell(0, 0).text == '任务')
entries = dict(zip(layout['titles'][2:], range(3, layout['pages'] + 1)))
for row in index.rows[1:]:
    for col in (0, 2):
        if row.cells[col].text:
            assert entries[row.cells[col].text] == int(row.cells[col + 1].text)
assert len(document.inline_shapes) == layout['figures'] + 1  # Cover logo is not a figure.
captions = [p.text for p in document.paragraphs if p.style.name == 'Caption']
assert len(captions) == layout['figures']
assert all(text.startswith(f'图 {n:02d} ') for n, text in enumerate(captions, 1))
with ZipFile(file) as archive:
    xml = archive.read('word/document.xml').decode('utf8')
    assert '0.7.2' not in xml
    assert 'fixture-key' not in xml
    assert '右侧显示路径、子文件夹' not in xml
report = {
    'result': 'PASS',
    'version': layout['version'],
    'pages': len(reader.pages),
    'annotatedFigures': layout['figures'],
    'pageTitlesAndIndexVerified': True,
    'visuallyReviewedPages': [],
    'renderer': json.loads((directory / 'rendered/renderer.json').read_text(encoding='utf8'))['method'],
    'docxSha256': hashlib.sha256(file.read_bytes()).hexdigest(),
    'pdfSha256': hashlib.sha256((directory / 'rendered/guide.pdf').read_bytes()).hexdigest(),
}
(root / 'test-results/guide-qa.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf8')
print(json.dumps(report, ensure_ascii=False, indent=2))
